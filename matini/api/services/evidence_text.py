"""Best-effort text extraction from uploaded evidence files (PDF, DOCX, TXT/MD)."""
import io
import os
from typing import Tuple

from matini.models.evidence import MAX_EXTRACTED_CHARS

NO_TEXT_ERROR = "ไม่พบข้อความในไฟล์ (อาจเป็นไฟล์สแกน)"
SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".txt", ".md")


def is_supported(filename: str) -> bool:
    return os.path.splitext(filename)[1].lower() in SUPPORTED_EXTENSIONS


def extract_text(filename: str, content_type: str, data: bytes) -> Tuple[str, str | None]:
    """Return (text truncated to MAX_EXTRACTED_CHARS, error message or None). Never raises."""
    ext = os.path.splitext(filename)[1].lower()
    try:
        if ext == ".pdf" or content_type == "application/pdf":
            text = _pdf_text(data)
        elif ext == ".docx":
            text = _docx_text(data)
        else:
            text = data.decode("utf-8-sig")
    except Exception as exc:  # corrupt or unsupported file: keep the upload, report why
        return "", f"อ่านข้อความจากไฟล์ไม่สำเร็จ: {exc.__class__.__name__}: {exc}"

    text = text.strip()
    if not text:
        return "", NO_TEXT_ERROR
    return text[:MAX_EXTRACTED_CHARS], None


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _docx_text(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(lines)
